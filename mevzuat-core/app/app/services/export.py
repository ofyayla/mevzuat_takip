"""Portal listesinin Excel (.xlsx) dışa aktarımı.

Liste ekranındaki süzgeç ve sıralamayla aynı kayıtlar, sayfalama olmadan tek sayfalık bir çalışma kitabına yazılır.
Kayıt başlıkları ve özetler dış kaynaklardan geldiği için ``=``, ``+``, ``-``, ``@`` ile başlayan metinler Excel'de
formül olarak çalışmasın diye başına kesme işareti konur (CSV/formül enjeksiyonu).
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_EXPORT_ROWS = 5000
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# (başlık, sütun genişliği, kayıttan değer)
COLUMNS: list[tuple[str, int, callable]] = [
    ("Kurum", 28, lambda r: r["issuer"]),
    ("Başlık", 60, lambda r: r["title"]),
    ("Tür", 18, lambda r: r["type"]),
    ("Yayım tarihi", 14, lambda r: r["publishDate"]),
    ("Yürürlük tarihi", 16, lambda r: r["effectiveDate"]),
    ("Önem", 12, lambda r: r["severity"]),
    ("Durum", 14, lambda r: r["status"]),
    ("Önerilen birimler", 36, lambda r: "; ".join(u["unit"] for u in r["suggestedUnits"])),
    ("Özet", 80, lambda r: r["summary"]),
    ("Konular", 40, lambda r: "; ".join(r["topics"])),
    ("Kaynak", 22, lambda r: ", ".join(r["sources"])),
    ("Bağlantı", 50, lambda r: r["sourceUrl"]),
    ("Karar veren", 22, lambda r: (r["decidedBy"] or {}).get("name")),
]


def _cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def build_workbook(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mevzuat"
    ws.append([name for name, _, _ in COLUMNS])
    for item in items:
        ws.append([_cell(get(item)) for _, _, get in COLUMNS])
    header_fill = PatternFill("solid", start_color="0B625B")
    for idx, (_, width, _) in enumerate(COLUMNS, start=1):
        head = ws.cell(row=1, column=idx)
        head.font, head.fill = Font(bold=True, color="FFFFFF"), header_fill
        head.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_filename(today: date) -> str:
    return f"mevzuat-{today.isoformat()}.xlsx"
