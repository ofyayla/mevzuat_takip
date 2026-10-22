"""Başkanlık Excel listesi ↔ sistem kayıtları karşılaştırması (salt okunur; veritabanına yazmaz).

Başkanlığın bankaya dağıttığı mevzuat listesi (Referans No, Konu, Talep Tipi, Gönderen/Düzenleyici Kurum, Risk
Seviyesi, Alıcılar, Oluşturulma Tarihi, Metin) yüklenir. Önce kurumla kapsamlanmış belge kimliği doğrulanır;
belirsiz kimlikler incelemeye ayrılır. Kimlik tanınmazsa İK-8 eşleştiricisi (``parallel.match_detection``) ve
ham belge metinlerinde bulanık arama kullanılır.

İki yön ölçülür:
  Excel → sistem   relevant (YZ ilgili dedi) / irrelevant (YZ eledi) / bulunamadı (nedeniyle) — recall
  sistem → Excel   dönemde portalda görünen ama listede olmayan kayıtlar (fazla uyarı) — precision

"Oluşturulma Tarihi" yayım tarihi değil, listeye giriş tarihidir (BDDK 30.09'da yayımlar, listeye 01.10'da girer);
bu yüzden eşleştirme penceresi geriye doğru geniştir.
"""
from __future__ import annotations

import html
import re
from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz
from selectolax.parser import HTMLParser
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.collectors.dates import parse_tr_date, tr_lower
from app.db import ManualDetection, RawDocument, Regulation
from app.evaluation.parallel import (
    TITLE_THRESHOLD,
    classify_detection,
    issuer_code,
    match_detection,
    read_records,
)
from app.evaluation.references import ReferenceResult, match_reference
from app.processing.normalize import ISSUERS, issuers_compatible, match_issuer, normalize_text, title_key
from app.services.export import _cell
from app.services.regulations import VISIBLE_STATUSES

WINDOW = (timedelta(days=14), timedelta(days=2))   # (geriye, ileriye): kayıt tarihi yayımdan sonra gelir
TEXT_THRESHOLD = 65
TEXT_NEEDLE_CHARS = 400        # Excel metninin aranan baş kısmı
TEXT_HAYSTACK_CHARS = 20000    # ham belgede bakılan kısım (uzun PDF'lerde süre sınırı)
EXTRA_LEAD = timedelta(days=3)  # ters yön dönemi: ilk kayıt tarihinden bu kadar önce yayımlananlar da sayılır
GENERIC_TITLE_TOKENS = 4        # bu kadar ya da daha az sözcüklü sistem başlığı "genel" sayılır (metinle doğrulanır)

COLUMNS = {"referans no": "ref", "referans": "ref", "konu": "title", "başlık": "title", "baslik": "title",
           "talep tipi": "request_type", "gönderen/düzenleyici kurum": "issuer", "düzenleyici kurum": "issuer",
           "kurum": "issuer", "risk seviyesi": "severity", "önem": "severity", "alıcılar": "recipients",
           "oluşturulma tarihi": "created", "tarih": "created", "yayım tarihi": "created", "metin": "text",
           "url": "url", "bağlantı": "url"}

GROUPS = {"detected": "relevant", "below_threshold": "irrelevant"}   # geri kalan kategoriler: missing
LABELS = {"detected": "İlgili (yakalandı)", "below_threshold": "İlgisiz (YZ eledi)", "not_collected": "Toplanmadı",
          "not_processed": "İşlenmedi", "baseline": "İlk tarama (YZ'ye gitmedi)",
          "dedupe_lost": "Birleştirmede kayboldu", "reference_review": "Belge kimliği: inceleme gerekli",
          "reference_raw_only": "Ham belge bulundu; sistem kaydı yok"}
METHODS = {"url": "Bağlantı", "exact_key": "Karar/sıra no", "fuzzy_title": "Başlık benzerliği",
           "fuzzy_title_raw": "Başlık (ham belge)", "fuzzy_text": "Metin benzerliği",
           "exact_reference": "Kurum + belge numarası"}


# ------------------------------------------------------------------------------------------------ giriş


def _header(h: str) -> str:
    return re.sub(r"\s*/\s*", "/", re.sub(r"\s+", " ", tr_lower(str(h or "").strip())))


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return parse_tr_date(str(v)) if v not in (None, "") else None


def clean_body(raw) -> str:
    """Metin sütunu: HTML etiketleri ve varlıkları (``&nbsp;``) temizlenir. "*****" ayracı varsa kurumun asıl
    duyuru metni ayraçtan sonradır; önündeki "… ektedir" giriş cümlesi atılır."""
    t = str(raw or "")
    if re.search(r"<[a-zA-Z/!]", t):
        tree = HTMLParser(t)
        node = tree.body or tree.root
        t = node.text(separator=" ") if node is not None else ""
    t = html.unescape(t)
    parts = re.split(r"\*{5,}", t, maxsplit=1)
    if len(parts) == 2 and len(parts[1].strip()) > 50:
        t = parts[1]
    return normalize_text(t)


def parse_sheet(filename: str, content: bytes) -> list[dict]:
    out = []
    for r in read_records(filename, content):
        row = {}
        for k, v in r.items():
            key = COLUMNS.get(_header(k))
            if key and key not in row:
                row[key] = v
        title = str(row.get("title") or "").strip()
        text = clean_body(row.get("text"))
        if not title and not text:
            continue
        out.append({"ref": str(row.get("ref") or "").strip() or None, "title": title or text[:200],
                    "request_type": str(row.get("request_type") or "").strip() or None,
                    "issuer": str(row.get("issuer") or "").strip() or None,
                    "severity": str(row.get("severity") or "").strip() or None,
                    "recipients": str(row.get("recipients") or "").strip() or None,
                    "created": _date(row.get("created")), "text": text,
                    "url": str(row.get("url") or "").strip() or None})
    return out


# ------------------------------------------------------------------------------------------------ eşleştirme


def _fold(text: str) -> str:
    return re.sub(r"[\W_]+", " ", tr_lower(text or "")).strip()


def _match_text(session: Session, row: dict, issuer: str | None, threshold: float) -> tuple[RawDocument | None, float]:
    """Excel metninin baş kısmı, pencere içindeki ham belge metinlerinde (ekler dahil) bulanık aranır."""
    needle = _fold(row["text"])[:TEXT_NEEDLE_CHARS]
    if len(needle) < 60 or row["created"] is None:
        return None, 0.0
    d = row["created"]
    q = select(RawDocument).where(RawDocument.is_latest.is_(True), RawDocument.role != "listing",
                                  RawDocument.text.is_not(None),
                                  RawDocument.published_at.between(d - WINDOW[0], d + WINDOW[1]))
    best, best_score = None, 0.0
    for raw in session.scalars(q):
        src = raw.source_code if raw.source_code in ISSUERS else None
        if issuer and src and not issuers_compatible(issuer, src):
            continue
        sc = fuzz.partial_ratio(needle, _fold(raw.text[:TEXT_HAYSTACK_CHARS]), score_cutoff=threshold)
        # eşit skorda düzenlemeye bağlı belge tercih edilir
        if sc > best_score or (sc == best_score and best is not None and best.regulation_id is None
                               and raw.regulation_id is not None):
            best, best_score = raw, sc
    return (best, best_score) if best is not None and best_score >= threshold else (None, 0.0)


def match_row(session: Session, row: dict, *, title_threshold: float = TITLE_THRESHOLD,
              text_threshold: float = TEXT_THRESHOLD) -> tuple[ManualDetection, ReferenceResult | None]:
    """Oturuma eklenmeyen (kalıcı olmayan) bir ManualDetection üzerinde eşleştirir."""
    md = ManualDetection(title=row["title"], issuer=row["issuer"], publish_date=row["created"], url=row["url"])
    if md.url:
        match_detection(session, md, window=WINDOW, title_threshold=title_threshold)
        if md.match_method == "url":
            return md, None
    reference = match_reference(session, row, WINDOW)
    if reference is not None:
        md.matched_raw_document_id = md.matched_regulation_id = None
        md.match_method = md.match_score = None
        if reference.raw is not None:
            md.matched_raw_document_id, md.matched_regulation_id = reference.raw.id, reference.raw.regulation_id
            md.match_method = "exact_reference"
        return md, reference
    match_detection(session, md, window=WINDOW, title_threshold=title_threshold)
    if not row["text"]:
        return md, None
    issuer = issuer_code(row["issuer"]) or match_issuer(row["title"])
    if md.matched_regulation_id is None and md.matched_raw_document_id is None:
        raw, score = _match_text(session, row, issuer, text_threshold)
        if raw is not None:
            md.matched_raw_document_id, md.matched_regulation_id = raw.id, raw.regulation_id
            md.match_method, md.match_score = "fuzzy_text", round(score / 100, 3)
    elif md.match_method == "fuzzy_title":
        # "Basın Duyurusu" gibi genel başlık, içinde bu sözcükler geçen her Konu'ya token_set ile 100 verir;
        # pencerede birden çok aday olabileceği için metin benzerliği farklı bir kaydı gösteriyorsa o tercih edilir
        reg = session.get(Regulation, md.matched_regulation_id)
        if reg is not None and len(title_key(reg.title).split()) <= GENERIC_TITLE_TOKENS:
            raw, score = _match_text(session, row, issuer, text_threshold)
            if raw is not None and raw.regulation_id and raw.regulation_id != md.matched_regulation_id:
                md.matched_raw_document_id, md.matched_regulation_id = raw.id, raw.regulation_id
                md.match_method, md.match_score = "fuzzy_text", round(score / 100, 3)
    return md, None


# ------------------------------------------------------------------------------------------------ karşılaştırma


def _reg_out(reg: Regulation | None) -> dict | None:
    if reg is None:
        return None
    rel = (reg.classification or {}).get("relevance") or {}
    return {"id": reg.id, "title": reg.title, "issuer": reg.issuer_name or reg.issuer,
            "publishDate": reg.publish_date.isoformat() if reg.publish_date else None,
            "status": reg.processing_status, "isRelevant": reg.is_relevant, "relevanceScore": reg.relevance_score,
            "severity": reg.severity, "rationale": rel.get("rationale")}


def compare(session: Session, rows: list[dict], *, title_threshold: float = TITLE_THRESHOLD,
            text_threshold: float = TEXT_THRESHOLD, date_from: date | None = None,
            date_to: date | None = None) -> dict:
    items, matched_ids = [], set()
    by_issuer: dict[str, Counter] = {}
    sev_pairs = []
    for i, row in enumerate(rows, start=1):
        md, reference = match_row(session, row, title_threshold=title_threshold, text_threshold=text_threshold)
        cat, reg = classify_detection(session, md)
        if reference is not None and reference.status != "matched":
            cat = "reference_review"
        elif reference is not None and reference.raw is not None and reg is None:
            cat = "reference_raw_only"
        group = GROUPS.get(cat, "missing")
        if reg is not None:
            matched_ids.add(reg.id)
        if reg is not None and row["severity"] and reg.severity:
            sev_pairs.append((row["severity"], reg.severity))
        code = issuer_code(row["issuer"]) or row["issuer"] or "—"
        by_issuer.setdefault(code, Counter())[group] += 1
        items.append({"no": i, "ref": row["ref"], "title": row["title"], "issuer": row["issuer"],
                      "requestType": row["request_type"], "severity": row["severity"],
                      "recipients": row["recipients"],
                      "createdAt": row["created"].isoformat() if row["created"] else None,
                      "category": cat, "categoryLabel": LABELS.get(cat, cat), "group": group,
                      "method": md.match_method, "methodLabel": METHODS.get(md.match_method or "", None),
                      "score": md.match_score, "rawDocumentId": md.matched_raw_document_id,
                      "reference": reference.output() if reference else None,
                      "regulation": _reg_out(reg)})

    total = len(items)
    groups = Counter(it["group"] for it in items)
    reasons = Counter(it["category"] for it in items if it["group"] == "missing")

    # ters yön: dönemde portalda görünen ama listedeki hiçbir satırla eşleşmeyen kayıtlar
    dates = [r["created"] for r in rows if r["created"]]
    lo = date_from or (min(dates) - EXTRA_LEAD if dates else None)
    hi = date_to or (max(dates) if dates else None)
    shown, extra = [], []
    if lo and hi:
        shown = session.scalars(select(Regulation).where(
            Regulation.publish_date.between(lo, hi), Regulation.is_relevant.is_(True),
            Regulation.processing_status.in_(VISIBLE_STATUSES)).order_by(Regulation.publish_date.desc())).all()
        extra = [_reg_out(r) for r in shown if r.id not in matched_ids]
    in_list = len(shown) - len(extra)

    return {
        "summary": {"total": total, "relevant": groups["relevant"], "irrelevant": groups["irrelevant"],
                    "missing": groups["missing"], "missingByReason": dict(reasons),
                    "recall": round(groups["relevant"] / total, 3) if total else None,
                    "foundRate": round((groups["relevant"] + groups["irrelevant"]) / total, 3) if total else None},
        "precision": {"period": {"from": lo.isoformat() if lo else None, "to": hi.isoformat() if hi else None},
                      "shown": len(shown), "inList": in_list, "extra": len(extra),
                      "precision": round(in_list / len(shown), 3) if shown else None},
        "severityAgreement": {"n": len(sev_pairs),
                              "exact": round(sum(a == b for a, b in sev_pairs) / len(sev_pairs), 3)
                              if sev_pairs else None},
        "byIssuer": [{"issuer": k, "total": sum(v.values()), "relevant": v["relevant"],
                      "irrelevant": v["irrelevant"], "missing": v["missing"]}
                     for k, v in sorted(by_issuer.items(), key=lambda kv: -sum(kv[1].values()))],
        "thresholds": {"title": title_threshold, "text": text_threshold},
        "items": items,
        "extra": extra,
    }


# ------------------------------------------------------------------------------------------------ Excel çıktısı

_HEADER_FILL = PatternFill("solid", start_color="0B625B")

ITEM_COLUMNS: list[tuple[str, int, callable]] = [
    ("Sıra", 6, lambda it: it["no"]),
    ("Referans No", 22, lambda it: it["ref"]),
    ("Konu", 60, lambda it: it["title"]),
    ("Kurum", 28, lambda it: it["issuer"]),
    ("Oluşturulma", 12, lambda it: it["createdAt"]),
    ("Risk seviyesi", 12, lambda it: it["severity"]),
    ("Sonuç", 20, lambda it: {"reference_review": "İnceleme gerekli", "reference_raw_only": "Toplandı"}.get(
        it["category"], {"relevant": "Relevant", "irrelevant": "Irrelevant"}.get(it["group"], "Bulunamadı"))),
    ("Ayrıntı", 26, lambda it: it["categoryLabel"]),
    ("Eşleşme yöntemi", 18, lambda it: it["methodLabel"]),
    ("Eşleşme skoru", 10, lambda it: it["score"]),
    ("Belge kimliği", 24, lambda it: (it.get("reference") or {}).get("number")),
    ("Karar tarihi", 14, lambda it: (it.get("reference") or {}).get("date")),
    ("Kimlik kontrolü", 18, lambda it: (it.get("reference") or {}).get("status")),
    ("Aday kayıtlar", 24, lambda it: ", ".join(map(str, (it.get("reference") or {}).get("candidateIds", [])))),
    ("Sistem kayıt no", 10, lambda it: (it["regulation"] or {}).get("id")),
    ("Ham belge no", 12, lambda it: it.get("rawDocumentId")),
    ("Sistem başlığı", 60, lambda it: (it["regulation"] or {}).get("title")),
    ("Sistem yayım tarihi", 12, lambda it: (it["regulation"] or {}).get("publishDate")),
    ("İlgililik skoru", 10, lambda it: (it["regulation"] or {}).get("relevanceScore")),
    ("Sistem önemi", 10, lambda it: (it["regulation"] or {}).get("severity")),
    ("YZ gerekçesi", 70, lambda it: (it["regulation"] or {}).get("rationale")),
]
EXTRA_COLUMNS: list[tuple[str, int, callable]] = [
    ("Kayıt no", 10, lambda r: r["id"]),
    ("Başlık", 70, lambda r: r["title"]),
    ("Kurum", 28, lambda r: r["issuer"]),
    ("Yayım tarihi", 12, lambda r: r["publishDate"]),
    ("İlgililik skoru", 10, lambda r: r["relevanceScore"]),
    ("Önem", 10, lambda r: r["severity"]),
    ("YZ gerekçesi", 70, lambda r: r["rationale"]),
]


def _table(ws, columns, rows) -> None:
    ws.append([name for name, _, _ in columns])
    for r in rows:
        ws.append([_cell(get(r)) for _, _, get in columns])
    for idx, (_, width, _) in enumerate(columns, start=1):
        head = ws.cell(row=1, column=idx)
        head.font, head.fill = Font(bold=True, color="FFFFFF"), _HEADER_FILL
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(result: dict) -> bytes:
    s, p = result["summary"], result["precision"]
    pct = lambda v: f"%{v * 100:.1f}" if v is not None else "—"   # noqa: E731
    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    summary = [("Excel satırı", s["total"]), ("Relevant (sistem buldu, YZ ilgili dedi)", s["relevant"]),
               ("Irrelevant (sistem buldu, YZ eledi)", s["irrelevant"]), ("Kesin eşleşme yok / inceleme", s["missing"])]
    summary += [(f"  · {LABELS.get(k, k)}", v) for k, v in s["missingByReason"].items()]
    summary += [("Recall (relevant / toplam)", pct(s["recall"])), ("Bulunma oranı", pct(s["foundRate"])),
                ("", None),
                ("Ters yön dönemi", f'{p["period"]["from"] or "—"} – {p["period"]["to"] or "—"}'),
                ("Portalda görünen kayıt", p["shown"]), ("Listede olan", p["inList"]),
                ("Fazla uyarı (listede yok)", p["extra"]), ("Precision", pct(p["precision"])),
                ("", None),
                ("Risk seviyesi uyumu", f'{pct(result["severityAgreement"]["exact"])} '
                                        f'({result["severityAgreement"]["n"]} kayıt)'),
                ("Başlık eşiği / metin eşiği", f'{result["thresholds"]["title"]} / {result["thresholds"]["text"]}')]
    for label, value in summary:
        ws.append([label, value])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 42, 24
    for c in ws["A"]:
        c.font = Font(bold=True)
    _table(wb.create_sheet("Excel satırları"), ITEM_COLUMNS, result["items"])
    _table(wb.create_sheet("Fazla uyarılar"), EXTRA_COLUMNS, result["extra"])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def compare_filename(today: date) -> str:
    return f"karsilastirma-{today.isoformat()}.xlsx"
