"""Publication date ordering is shared by inbox pagination and Excel exports."""
from __future__ import annotations

from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.db import Regulation, utcnow
from app.services.regulations import today
from tests.test_api import client  # noqa: F401
from tests.test_processing import env  # noqa: F401


@pytest.mark.parametrize("direction", ["asc", "desc"])
def test_publication_order_across_pages_and_export(settings, env, client, direction):  # noqa: F811
    c, sf = client
    older = today(settings) - timedelta(days=400)
    newer = today(settings) - timedelta(days=1)
    detected = utcnow()
    with sf() as session:
        records = [
            Regulation(title=title, canonical_key=f"publication:{title}", publish_date=published,
                       detected_at=detected, issuer="BDDK", primary_source_code="BDDK",
                       is_relevant=True, processing_status="READY", severity="Orta")
            for title, published in [("Older", older), ("Newer", newer), ("Undated", None), ("Older tie", older)]
        ]
        session.add_all(records)
        session.add(Regulation(title="Filtered out", canonical_key="publication:excluded", publish_date=newer,
                               detected_at=detected, primary_source_code="BDDK",
                               is_relevant=True, processing_status="READY", severity="Kritik"))
        session.flush()
        old_id, new_id, null_id, tie_id = [record.id for record in records]
        session.commit()

    expected = [tie_id, old_id, new_id, null_id] if direction == "asc" else [new_id, tie_id, old_id, null_id]
    query = {"queue": "tum", "sort": "publish_date", "direction": direction, "severity": "Orta"}
    items = []
    for page in (1, 2):
        response = c.get("/api/v1/regulations", params={**query, "page": page, "page_size": 2})
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 4
        items.extend(body["items"])
    assert [item["id"] for item in items] == expected
    assert items[-1]["publishDate"] is None
    assert {item["publishDate"] for item in items} == {older.isoformat(), newer.isoformat(), None}

    response = c.get("/api/v1/regulations/export", params=query)
    assert response.status_code == 200
    rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
    assert rows[0][3] == "Yayım tarihi"
    assert [row[1] for row in rows[1:]] == [item["title"] for item in items]
    assert [row[3] for row in rows[1:]] == [item["publishDate"] for item in items]
