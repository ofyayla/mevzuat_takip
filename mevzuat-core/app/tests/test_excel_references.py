# ruff: noqa: F811
import json
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.api.main import create_app
from app.db import ManualDetection, RawDocument, Regulation, RegulationKey, Source
from app.evaluation.excel_compare import build_workbook, compare
from app.evaluation.references import extract_reference
from app.processing.dedupe import canonical_url
from app.services.regulations import ApiError
from tests.test_excel_compare import _reg, _xlsx
from tests.test_processing import add_raw, env  # noqa: F401

DAY = date(2026, 10, 3)
DECISION = "Kurul Karar Organı’nın 03/10/2026 tarih ve 65/1779 sayılı Kararı Uyarınca Yapılan Duyuru"


def row(title=DECISION, text="", issuer="SPK", created=DAY, url=None):
    return dict(ref="test", title=title, text=text, issuer=issuer, created=created, url=url,
                request_type=None, severity=None, recipients=None)


def document(env, ext, title, text, *, issuer="SPK", pub=DAY, reg_id=None, extra=None):
    sf, storage = env
    with sf() as s:
        if s.get(Source, issuer) is None:
            s.add(Source(code=issuer, name=issuer))
            s.commit()
    if reg_id is None:
        reg_id = _reg(sf, ext, title, "READY", pub, issuer=issuer, relevant=True)
    raw_id = add_raw(sf, storage, source=issuer, channel="test", external_id=ext,
                     title=title, url=f"https://example.org/{ext}", published=pub, extra=extra)
    with sf() as s:
        raw = s.get(RawDocument, raw_id)
        raw.text, raw.regulation_id = text, reg_id
        s.commit()
    return reg_id, raw_id


@pytest.mark.parametrize(("title", "body", "issuer", "kind", "number"), [
    (DECISION, "", "SPK", "decision", "65/1779"),
    ("SPK Bülteni 2026/68", "", "SPK", "bulletin", "2026/68"),
    ("Rehber güncellemesi", "2026/54 sayılı SPK Bülteni yayımlandı.", "SPK", "bulletin", "2026/54"),
    ("Özelge", "E-62030549-120-1372289 sayılı özelge", "GIB", "ruling", "e-62030549-120-1372289"),
    ("İzin iptali", "12100/21699 sayılı ve 07.10.2026 tarihli TCMB kararı", "TCMB", "decision", "12100/21699"),
    ("İzin iptali", "12100/21699 sayı ve 07.10.2026 tarihli Parakolay hakkında TCMB kararı", "TCMB", "decision", "12100/21699"),
    ("İzin iptali", "10.10.2026 tarih ve 33396 sayılı Resmi Gazete'de yayımlanan 12100/21699 sayılı ve "
     "07.10.2026 tarihli Parakolay konulu TCMB kararı", "TCMB", "decision", "12100/21699"),
    ("Genel mektup", "2175 sayılı Genel Mektup ektedir.", "TAKASBANK", "letter", "2175"),
    ("Yangın Genelgesi", "2026/29 sayılı Genelge", "SEDDK", "circular", "2026/29"),
    ("İzin", "11537 sayılı ve Tasarruf Finansman Faaliyetleri Hk. Kurul Kararı", "BDDK", "decision", "11537"),
])
def test_typed_reference_extraction(title, body, issuer, kind, number):
    ref = extract_reference(title, body, issuer)
    assert ref is not None and (ref.kind, ref.number) == (kind, number)


def test_decision_inside_bulletin_overrides_fuzzy_thresholds(env):
    rid, raw_id = document(env, "bulletin", "SPK Bülteni 2026/68", DECISION + "\nÖnceki 57/1706 sayılı karar uzatıldı.")
    with env[0]() as s:
        result = compare(s, [row()], title_threshold=100, text_threshold=100)
    item = result["items"][0]
    assert item["regulation"]["id"] == rid
    assert item["rawDocumentId"] == raw_id
    assert item["method"] == "exact_reference" and item["score"] is None
    assert item["reference"]["status"] == "matched"
    book = openpyxl.load_workbook(BytesIO(build_workbook(result)))
    values = list(book["Excel satırları"].values)
    assert values[1][values[0].index("Belge kimliği")] == "65/1779"
    assert values[1][values[0].index("Karar tarihi")] == "2026-10-03"
    assert values[1][values[0].index("Ham belge no")] == raw_id


def test_same_regulation_copies_are_not_ambiguous_but_two_records_are(env):
    rid, _ = document(env, "one", DECISION, DECISION)
    document(env, "copy", DECISION, DECISION, reg_id=rid)
    with env[0]() as s:
        assert compare(s, [row()])["items"][0]["reference"]["status"] == "matched"
    other, _ = document(env, "two", DECISION, DECISION)
    with env[0]() as s:
        item = compare(s, [row()])["items"][0]
    assert item["regulation"] is None and item["category"] == "reference_review"
    assert item["reference"]["status"] == "ambiguous"
    assert item["reference"]["candidateIds"] == sorted([rid, other])


def test_citation_is_not_a_primary_decision(env):
    rid, _ = document(env, "citation", "SPK Bülteni 2026/68",
                      DECISION.replace("65/1779", "65/1780") + "\n65/1779 sayılı karara atıf yapılmıştır.")
    with env[0]() as s:
        item = compare(s, [row()], title_threshold=50, text_threshold=50)["items"][0]
    assert item["regulation"] is None and item["category"] == "reference_review"
    assert item["reference"]["status"] == "citation" and item["reference"]["candidateIds"] == [rid]


@pytest.mark.parametrize("change", ["issuer", "date", "number", "old_version", "listing"])
def test_conflicting_identifiers_do_not_fall_back_to_similar_titles(env, change):
    title = DECISION.replace("1779", "1778") if change == "number" else DECISION
    text = title.replace("03/10/2026", "02/10/2026") if change == "date" else title
    if change == "date":
        title = "Benzer konulu duyuru"
    _, raw_id = document(env, "wrong", title, text, issuer="BDDK" if change == "issuer" else "SPK")
    with env[0]() as s:
        raw = s.get(RawDocument, raw_id)
        if change == "old_version":
            raw.is_latest = False
        if change == "listing":
            raw.role = "listing"
        s.commit()
        item = compare(s, [row()], title_threshold=50, text_threshold=50)["items"][0]
    assert item["regulation"] is None
    assert item["category"] == "reference_review"


def test_reference_only_in_excel_citation_is_not_automatic(env):
    document(env, "primary", DECISION, DECISION)
    with env[0]() as s:
        item = compare(s, [row("Konu hakkında bilgi", "Önceki karar: " + DECISION)])["items"][0]
    assert item["regulation"] is None and item["reference"]["status"] == "citation"


def test_gib_number_from_metadata(env):
    rid, _ = document(env, "gib", "Vergi istisnası", "Yerleşme şartının uygulanması hakkında açıklama.",
                      issuer="GIB", extra={"ozelge_no": "E-62030549-120-1372289"})
    with env[0]() as s:
        item = compare(s, [row("İstisna", "E-62030549-120-1372289 sayılı özelge", "GIB")])["items"][0]
    assert item["regulation"]["id"] == rid


def test_tcmb_reversed_number_and_date(env):
    rid, _ = document(env, "tcmb", "Parakolay faaliyet izni",
                      "Karar Sayısı: 12100/21699\nKarar Tarihi: 7/10/2026\nParakolay", issuer="TCMB",
                      pub=date(2026, 10, 10))
    with env[0]() as s:
        item = compare(s, [row("İzin iptali", "12100/21699 sayılı ve 07.10.2026 tarihli TCMB kararı",
                               "TCMB", date(2026, 10, 12))])["items"][0]
    assert item["regulation"]["id"] == rid


def test_missing_context_is_reviewed(env):
    document(env, "primary", DECISION, DECISION)
    with env[0]() as s:
        # Tarihi açıkça verilen karar, Excel kayıt tarihi boş olsa da doğrulanabilir.
        assert compare(s, [row(created=None)])["items"][0]["reference"]["status"] == "matched"
        item = compare(s, [row("SPK Bülteni 2026/68", created=None)])["items"][0]
    assert item["reference"]["status"] == "unverified"


def test_gazette_number_is_not_a_decision():
    assert extract_reference("Duyuru", "27.08.2015 tarih ve 29458 sayılı Resmi Gazete", "GIB") is None


def test_primary_header_wins_over_cited_older_decision():
    ref = extract_reference(
        "Faaliyet izni iptali",
        "Karar Sayısı: 12100/21699\nKarar Tarihi: 7/10/2026\n"
        "28/7/2022 tarih ve 11285/20884 sayılı Kararı ile verilen iznin iptaline karar verilmiştir.",
        "TCMB",
    )
    assert ref and ref.number == "12100/21699" and ref.dated == date(2026, 10, 7)


def test_late_publication_uses_excel_window_and_exact_decision_date(env):
    title = "16/09/2026 tarihli ve 2026/2035 sayılı Kararı"
    rid, _ = document(env, "late", title, "", pub=date(2026, 10, 8), issuer="KVKK")
    with env[0]() as s:
        item = compare(s, [row(title=title, issuer="KVKK", created=date(2026, 10, 8))])["items"][0]
    assert item["regulation"]["id"] == rid


@pytest.mark.parametrize("change", ["number", "date"])
def test_title_cannot_override_conflicting_structured_header(env, change):
    number = "65/1778" if change == "number" else "65/1779"
    day = "02" if change == "date" else "03"
    document(env, "conflict", DECISION, f"Karar No: {number}\nKarar Tarihi: {day}/10/2026")
    with env[0]() as s:
        item = compare(s, [row()])["items"][0]
    assert item["category"] == "reference_review" and item["reference"]["status"] == "conflict"


def test_cited_institution_cannot_override_known_publisher(env):
    document(env, "issuer-conflict", DECISION, "Sermaye Piyasası Kurulu\n" + DECISION, issuer="BDDK")
    with env[0]() as s:
        assert compare(s, [row()])["items"][0]["category"] == "reference_review"


def test_raw_without_regulation_is_collected_not_missing_crawl(env):
    rid, raw_id = document(env, "unlinked", DECISION, DECISION)
    with env[0]() as s:
        s.get(RawDocument, raw_id).regulation_id = None
        s.commit()
        item = compare(s, [row()])["items"][0]
    assert item["regulation"] is None and item["rawDocumentId"] == raw_id
    assert item["category"] == "reference_raw_only"


def test_reference_number_substrings_are_not_equal(env):
    document(env, "longer-number", DECISION.replace("1779", "17790"), "")
    with env[0]() as s:
        item = compare(s, [row()])["items"][0]
    assert item["regulation"] is None and item["reference"]["candidateIds"] == []


def test_unknown_url_cannot_leave_an_unverified_fuzzy_match(env):
    document(env, "wrong", DECISION.replace("1779", "1778"), "")
    with env[0]() as s:
        item = compare(s, [row(url="https://example.org/unknown")])["items"][0]
    assert item["category"] == "reference_review"
    assert item["regulation"] is None and item["rawDocumentId"] is None and item["method"] is None


def test_explicit_url_keeps_existing_precedence(env):
    rid, _ = document(env, "url", DECISION.replace("1779", "1778"), "")
    url = "https://example.org/direct"
    with env[0]() as s:
        s.add(RegulationKey(key="url:" + canonical_url(url), regulation_id=rid))
        s.commit()
        item = compare(s, [row(url=url)])["items"][0]
    assert item["regulation"]["id"] == rid and item["method"] == "url" and item["reference"] is None


def test_law_number_in_header_is_not_a_primary_decision():
    assert extract_reference(
        "Duyuru", "20/6/2013 tarih ve 6493 sayılı Ödeme ve Menkul Kıymet Mutabakat Sistemleri Hakkında Kanun",
        "TCMB",
    ) is None


def test_undated_number_does_not_match_prior_year(env):
    document(env, "prior-year", "Karar No: 11537", "", issuer="BDDK", pub=date(2025, 12, 31))
    with env[0]() as s:
        item = compare(s, [row("Karar No: 11537", issuer="BDDK", created=date(2026, 1, 2))])["items"][0]
    assert item["category"] == "reference_review"


def test_review_accept_and_undo_preserve_ai_label_and_metrics(env):
    first, _ = document(env, "candidate-a", DECISION, DECISION)
    other, _ = document(env, "candidate-b", DECISION, DECISION)
    with env[0]() as s:
        r = s.get(Regulation, other)
        r.processing_status, r.is_relevant, r.severity = "IRRELEVANT", False, "Orta"
        s.commit()
        before = compare(s, [row()])
        assert before["summary"]["review"] == 1 and before["summary"]["notFound"] == 0
        assert before["items"][0]["matchState"] == "review"
        assert {c["regulation"]["id"] for c in before["items"][0]["candidates"]} == {first, other}
        assert before["items"][0]["candidates"][0]["textExcerpt"] == DECISION
        other_candidate = next(c for c in before["items"][0]["candidates"] if c["regulation"]["id"] == other)
        assert other_candidate["selection"]["group"] == "irrelevant"
        assert other_candidate["selection"]["regulation"]["id"] == other
        accepted = compare(s, [row()], selections={1: other})
        item = accepted["items"][0]
        assert item["method"] == "manual_selection" and item["score"] is None
        assert item["group"] == "irrelevant" and item["selectedRegulationId"] == other
        assert item["group"] == other_candidate["selection"]["group"]
        assert item["regulation"] == other_candidate["selection"]["regulation"]
        assert accepted["summary"]["review"] == 0 and accepted["summary"]["matched"] == 1
        assert compare(s, [row()])["summary"]["review"] == 1
        assert s.scalar(select(func.count()).select_from(ManualDetection)) == 0
        assert s.get(Regulation, other).is_relevant is False
        book = openpyxl.load_workbook(BytesIO(build_workbook(accepted)))
        values = list(book["Excel satırları"].values)
        assert values[1][values[0].index("Kullanıcı onaylı eşleşme")] == other


def test_missing_baseline_and_review_have_separate_counters(env):
    rid, _ = document(env, "baseline", DECISION, DECISION)
    with env[0]() as s:
        s.get(Regulation, rid).processing_status = "BASELINE"
        s.commit()
        result = compare(s, [row(), row(title="SPK Bülteni 2026/99")])
    assert result["summary"]["notFound"] == 1
    assert result["summary"]["review"] == 0 and result["summary"]["unprocessed"] == 1
    assert result["summary"]["matched"] == 1 and result["summary"]["matchRate"] == 0.5
    book = openpyxl.load_workbook(BytesIO(build_workbook(result)))
    values = list(book["Excel satırları"].values)
    assert values[1][values[0].index("Sonuç")] == "Bulundu; YZ sonucu yok"


def test_unverified_exact_title_is_review_candidate_not_automatic_match(env):
    title = "Kurumun uygulamaya ilişkin açıklaması"
    rid, _ = document(env, "suggestion", title, "Belge numarası metinde yok.")
    with env[0]() as s:
        result = compare(s, [row(title=title, text=DECISION)])
        item = result["items"][0]
        assert item["regulation"] is None and item["reference"]["status"] == "unverified"
        assert item["candidates"][0]["regulation"]["id"] == rid
        assert item["candidates"][0]["evidence"] == "similarity"
        assert result["summary"]["review"] == 1
        book = openpyxl.load_workbook(BytesIO(build_workbook(result)))
        values = list(book["Excel satırları"].values)
        assert values[1][values[0].index("Aday kayıtlar")] == str(rid)


@pytest.mark.parametrize("selections", [{1: 999999}, {2: 1}])
def test_invalid_manual_selection_is_rejected(env, selections):
    document(env, "candidate", DECISION, DECISION)
    with env[0]() as s, pytest.raises(ApiError):
        compare(s, [row()], selections=selections)


def test_manual_selections_multipart_json_and_export_are_read_only(settings, env):
    first, _ = document(env, "a", DECISION, DECISION)
    document(env, "b", DECISION, DECISION)
    c = TestClient(create_app(settings.model_copy(update={"cors_origins": ""}), env[0]))
    content = _xlsx([["REF", DECISION, "Bilgi", "SPK", "Orta", "Uyum", DAY, DECISION]])
    files = {"file": ("test.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    selections = json.dumps([{"row": 1, "regulation_id": first}])
    url = "/api/v1/evaluation/excel-compare"
    result = c.post(url, files=files, data={"selections": selections})
    assert result.status_code == 200, result.text
    assert result.json()["items"][0]["method"] == "manual_selection"
    export = c.post(url, params={"format": "xlsx"}, files=files, data={"selections": selections})
    assert export.status_code == 200, export.text
    assert openpyxl.load_workbook(BytesIO(export.content))["Excel satırları"].max_row == 2
    for invalid in ["not-json", '[{"row":1,"regulation_id":true}]']:
        assert c.post(url, files=files, data={"selections": invalid}).status_code == 422
    duplicated = json.dumps([{"row": 1, "regulation_id": first}] * 2)
    assert c.post(url, files=files, data={"selections": duplicated}).status_code == 400
    with env[0]() as s:
        assert s.scalar(select(func.count()).select_from(ManualDetection)) == 0
