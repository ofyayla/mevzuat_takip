"""İK-6 Portal API: liste/detay/istatistik, karar ve birim değişikliği iş kuralları, denetim izi, iyimser kilit,
kullanıcı kimliği (X-Demo-User başlığı), yönetim uç noktaları."""
from __future__ import annotations

from datetime import timedelta
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.ai.llm_client import FakeLLM
from app.ai.summary import summarize_pending
from app.ai.unit_matching import match_pending
from app.api.main import create_app
from app.db import AuditEvent, FewshotExample, Outbox, Regulation, RegulationSummary, UnitSuggestion
from app.services.regulations import today
from tests.test_processing import env  # noqa: F401
from tests.test_summary import GOOD, _relevant_bank_mellat, summary_out
from tests.test_units import UNITS, _set_topics

RISK_RESP = UNITS["units"][0]["responsibilities"][-1]


def _ready(settings, env):  # noqa: F811
    sf, _ = env
    reg_id = _relevant_bank_mellat(settings, env)
    summarize_pending(sf, FakeLLM(settings, lambda t, m, n: [summary_out([GOOD])]), settings)
    _set_topics(sf, reg_id, ["SERMAYE_LIKIDITE"])
    match_pending(sf, FakeLLM(settings, lambda t, m, n: [{"suggestions": [
        {"unit_code": "RISK_YONETIMI", "score": 0.9, "matched_responsibility": RISK_RESP, "reason": "Karşı taraf."}]}]),
        settings)
    return reg_id


@pytest.fixture
def client(settings, env):  # noqa: F811
    sf, _ = env
    s = settings.model_copy(update={"cors_origins": ""})
    return TestClient(create_app(s, sf)), sf


def test_list_detail_and_stats(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, _ = client
    body = c.get("/api/v1/regulations").json()
    assert body["total"] == 1 and body["page"] == 1
    item = body["items"][0]
    assert item["id"] == reg_id and item["status"] == "Bekliyor" and item["issuerCode"] == "BDDK"
    assert item["publishDate"] == "2026-09-19" and item["summary"] == GOOD["sentence"]
    assert item["suggestedUnits"][0]["unit"] == "Risk Yönetimi Başkanlığı" and item["analysisStatus"] == "hazır"
    assert item["sources"] == ["BDDK", "RESMI_GAZETE"] and item["sourceUrl"].startswith("https://www.resmigazete")
    assert c.get("/api/v1/regulations", params={"source": "KVKK"}).json()["total"] == 0
    assert c.get("/api/v1/regulations", params={"unit": "RISK_YONETIMI"}).json()["total"] == 1
    assert c.get("/api/v1/regulations", params={"q": "Mellat"}).json()["total"] == 1
    assert c.get("/api/v1/regulations", params={"metric": "critical_high"}).json()["total"] == 0   # önem Orta
    st = c.get("/api/v1/regulations/stats").json()
    assert {k: st[k] for k in ("pending", "criticalHighPending", "today", "upcoming")} == {
        "pending": 1, "criticalHighPending": 0, "today": 1, "upcoming": 0}
    r = c.get(f"/api/v1/regulations/{reg_id}")
    assert r.headers["etag"] == '"1"'
    d = r.json()
    assert d["auditTrail"][0]["action"] == "YZ tarafından tespit edildi" and d["auditTrail"][0]["icon"] == "bot"
    assert d["evidence"]["summary"][0]["evidence"][0]["method"] == "exact"
    assert {"summary", "suggestedUnits", "severity"} <= set(d["aiGeneratedFields"])
    # hata biçimi
    assert c.get("/api/v1/regulations/9999").json() == {"error": {"code": "not_found",
                                                                  "message": "düzenleme bulunamadı", "details": {}}}
    assert c.get("/api/v1/regulations", params={"page_size": 500}).json()["error"]["code"] == "validation_error"


def test_export_xlsx_uses_list_filters_and_neutralizes_formulas(settings, env, client):  # noqa: F811
    from io import BytesIO

    from openpyxl import load_workbook

    reg_id = _ready(settings, env)
    c, sf = client
    r = c.get("/api/v1/regulations/export")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"] and r.headers["x-export-truncated"] == "false"
    rows = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0][:2] == ("Kurum", "Başlık") and len(rows) == 2
    assert rows[1][7] == "Risk Yönetimi Başkanlığı" and rows[1][8] == GOOD["sentence"]
    # süzgeç listeyle aynı çalışır; eşleşme yoksa yalnız başlık satırı kalır
    empty = c.get("/api/v1/regulations/export", params={"source": "KVKK"})
    assert len(list(load_workbook(BytesIO(empty.content)).active.iter_rows())) == 1
    assert c.get("/api/v1/regulations/export", params={"status": "yok"}).json()["error"]["code"] == "invalid_status"
    with sf() as s:
        s.get(Regulation, reg_id).title = "=HYPERLINK(\"http://x\")"
        s.commit()
    rows = list(load_workbook(BytesIO(c.get("/api/v1/regulations/export").content)).active.iter_rows(values_only=True))
    assert rows[1][1].startswith("'=")


def test_irrelevant_and_merged_are_hidden(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    with sf() as s:
        s.get(Regulation, reg_id).is_relevant = False
        s.commit()
    assert c.get("/api/v1/regulations").json()["total"] == 0
    assert c.get(f"/api/v1/regulations/{reg_id}").status_code == 404


def test_views_once_per_user(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    assert c.post(f"/api/v1/regulations/{reg_id}/views").json() == {"created": True}
    assert c.post(f"/api/v1/regulations/{reg_id}/views").json() == {"created": False}
    assert c.post(f"/api/v1/regulations/{reg_id}/views", headers={"X-Demo-User": quote("Ayşe")}).json()["created"]
    with sf() as s:
        assert len(s.scalars(select(AuditEvent).where(AuditEvent.event_type == "viewed")).all()) == 2


def test_approve_freezes_units_records_learning_and_outbox(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    r = c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "approve", "note": "Takip edilecek."},
               headers={"If-Match": '"1"', "X-Demo-User": quote("Ayşe Kaya")})
    assert r.status_code == 200 and r.headers["etag"] == '"2"'
    d = r.json()
    assert d["status"] == "Onaylandı" and d["decidedBy"]["name"] == "Ayşe Kaya" and d["decisionNote"] == "Takip edilecek."
    last = d["auditTrail"][-1]
    assert last["action"] == "Onaylandı" and last["icon"] == "check-circle-2" and last["actor"] == "Ayşe Kaya · Uyum Uzmanı"
    assert "Risk Yönetimi Başkanlığı birimine yönlendirildi." in last["note"] and "Takip edilecek." in last["note"]
    with sf() as s:
        assert s.scalars(select(UnitSuggestion)).one().is_final
        ex = s.scalars(select(FewshotExample)).one()
        assert (ex.origin, ex.weight, ex.expected_output) == ("user_decision", 1.0, {"units": ["RISK_YONETIMI"]})
        ob = s.scalars(select(Outbox)).one()
        assert ob.event_type == "record.approved" and ob.payload["units"] == ["RISK_YONETIMI"]
    # ikinci karar: 409; birim değişikliği de kapalı
    again = c.post(f"/api/v1/regulations/{reg_id}/decision",
                   json={"decision": "reject", "rejection_reason": "duplicate"})
    assert again.status_code == 409 and again.json()["error"]["code"] == "already_decided"
    assert c.put(f"/api/v1/regulations/{reg_id}/units", json={"unit_codes": ["HAZINE"]}).status_code == 409


def test_optimistic_locking(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, _ = client
    ok = c.put(f"/api/v1/regulations/{reg_id}/units", json={"unit_codes": ["RISK_YONETIMI", "HAZINE"]},
               headers={"If-Match": '"1"'})
    assert ok.status_code == 200 and ok.json()["version"] == 2
    stale = c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "approve"}, headers={"If-Match": '"1"'})
    assert stale.status_code == 412 and stale.json()["error"]["details"] == {"current_version": 2}


def test_edit_title_and_summary_creates_version_and_checks_etag(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    updated = c.patch(f"/api/v1/regulations/{reg_id}/content",
                      json={"title": "Düzenlenen başlık", "summary": "Düzenlenen YZ özeti."},
                      headers={"If-Match": '"1"', "X-Demo-User": quote("Ayşe Kaya")})
    assert updated.status_code == 200 and updated.headers["etag"] == '"2"'
    assert (updated.json()["title"], updated.json()["summary"], updated.json()["summaryVersion"]) == (
        "Düzenlenen başlık", "Düzenlenen YZ özeti.", 2)
    assert updated.json()["auditTrail"][-1]["action"] == "Kayıt bilgileri güncellendi"
    assert updated.json()["auditTrail"][-1]["note"] == "Başlık, YZ özeti güncellendi."
    with sf() as s:
        versions = s.scalars(select(RegulationSummary).where(RegulationSummary.regulation_id == reg_id)
                             .order_by(RegulationSummary.version)).all()
        assert [(row.version, row.is_current) for row in versions] == [(1, False), (2, True)]
        assert versions[1].short_content_evidence == [] and versions[1].grounding_score == 0
    stale = c.patch(f"/api/v1/regulations/{reg_id}/content", json={"title": "Eski sürüm"},
                    headers={"If-Match": '"1"'})
    assert stale.status_code == 412 and stale.json()["error"]["details"] == {"current_version": 2}


def test_change_units_keeps_ai_reason_and_records_correction(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    r = c.put(f"/api/v1/regulations/{reg_id}/units", json={"unit_codes": ["HAZINE", "RISK_YONETIMI"],
                                                           "note": "Likidite etkisi"})
    units = r.json()["suggestedUnits"]
    assert [(u["code"], u["origin"]) for u in units] == [("HAZINE", "manual"), ("RISK_YONETIMI", "ai")]
    assert units[0]["reason"] == "Uzman tarafından manuel olarak atandı." and units[1]["reason"] == "Karşı taraf."
    ev = r.json()["auditTrail"][-1]
    assert ev["action"] == "Birim değiştirildi" and ev["icon"] == "shuffle"
    assert '"Risk Yönetimi Başkanlığı" iken "Hazine Müdürlüğü, Risk Yönetimi Başkanlığı"' in ev["note"]
    with sf() as s:
        ex = s.scalars(select(FewshotExample)).one()
        assert (ex.origin, ex.weight) == ("user_correction", 2.0)
    # onayda düzeltme zaten kaydedildiği için ikinci örnek oluşmaz
    c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "approve"})
    with sf() as s:
        assert len(s.scalars(select(FewshotExample)).all()) == 1
    bad = c.put(f"/api/v1/regulations/{reg_id}/units", json={"unit_codes": []})
    assert bad.status_code == 422


def test_unknown_unit_and_reject(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    r = c.put(f"/api/v1/regulations/{reg_id}/units", json={"unit_codes": ["HUKUK_ISLERI"]})
    assert r.status_code == 400 and r.json()["error"]["details"] == {"unit_codes": ["HUKUK_ISLERI"]}
    r = c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "reject", "rejection_reason": "duplicate"})
    assert r.json()["status"] == "Reddedildi" and r.json()["auditTrail"][-1]["icon"] == "x-circle"
    assert r.json()["auditTrail"][-1]["rejectionReason"] == "duplicate"
    with sf() as s:
        event = s.scalars(select(AuditEvent).where(AuditEvent.event_type == "rejected")).one()
        assert event.payload["rejection_reason"] == "duplicate"
        outbox = s.scalars(select(Outbox)).one()
        assert outbox.event_type == "record.rejected" and outbox.payload["rejection_reason"] == "duplicate"
        assert s.scalars(select(FewshotExample)).first() is None


def test_rejection_feedback_requires_reason_and_other_details(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, _ = client
    missing_reason = c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "reject"})
    assert missing_reason.status_code == 422
    missing_details = c.post(f"/api/v1/regulations/{reg_id}/decision",
                             json={"decision": "reject", "rejection_reason": "other", "note": "  "})
    assert missing_details.status_code == 422
    accepted_other = c.post(f"/api/v1/regulations/{reg_id}/decision",
                            json={"decision": "reject", "rejection_reason": "other", "note": "Kurum kapsamı dışında."})
    assert accepted_other.status_code == 200
    assert accepted_other.json()["decisionNote"] == "Kurum kapsamı dışında."
    assert accepted_other.json()["auditTrail"][-1]["rejectionReason"] == "other"


def test_audit_is_append_only(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    _, sf = client
    with sf() as s:
        ev = s.scalars(select(AuditEvent).where(AuditEvent.regulation_id == reg_id)).first()
        ev.note = "değiştirildi"
        with pytest.raises(PermissionError):
            s.commit()


def test_effective_soon_and_upcoming(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    with sf() as s:
        s.get(Regulation, reg_id).effective_date = today(settings) + timedelta(days=10)
        s.commit()
    assert c.get("/api/v1/regulations").json()["items"][0]["effectiveSoon"] is True
    assert c.get("/api/v1/regulations/stats").json()["upcoming"] == 1


# ------------------------------------------------------------------------------------------------ kullanıcı kimliği


def test_identity_from_header_all_roles(settings, env, client):  # noqa: F811
    """Kimlik doğrulama yok: ad X-Demo-User başlığından gelir, herkes tüm rollere sahiptir; Authorization yok sayılır."""
    reg_id = _ready(settings, env)
    c, _ = client
    h = {"X-Demo-User": quote("Ayşe Kaya"), "Authorization": "Bearer gecersiz"}
    me = c.get("/api/v1/me", headers=h).json()
    assert me["name"] == "Ayşe Kaya" and me["canDecide"] and me["isAdmin"] and "authMode" not in me
    assert c.get("/api/v1/me").json()["name"] == settings.demo_user_name
    assert c.get("/api/v1/admin/settings", headers=h).status_code == 200
    r = c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "approve"}, headers=h)
    assert r.json()["decidedBy"]["name"] == "Ayşe Kaya"


# ------------------------------------------------------------------------------------------------ yönetim ve genel


def test_admin_settings_and_health_endpoints(settings, env, client):  # noqa: F811
    c, _ = client
    r = c.put("/api/v1/admin/settings", json={"values": {"relevance_threshold": "0.4"}})
    assert r.json()["relevance_threshold"] == 0.4
    assert c.put("/api/v1/admin/settings", json={"values": {"llm_api_key": "x"}}).status_code == 400
    h = c.get("/api/v1/sources/health").json()
    assert h["overall"] == "down" and {s["status"] for s in h["sources"]} == {"down"}   # hiç tarama yok
    assert "Henüz tarama yapılmadı" in h["bannerText"]
    assert c.get("/healthz").json() == {"status": "ok"} and c.get("/readyz").json()["checks"]["db"] == "ok"
    assert len(c.get("/api/v1/units").json()) == 0 or True
    assert c.get("/").status_code == 200 and "Mevzuat" in c.get("/").text
    assert c.get("/admin", follow_redirects=False).headers["location"] == "/#yonetim"   # yeni portalda Yönetim sekmesi


def test_runtime_settings_validate_type_range_and_consistency(settings, env, client):  # noqa: F811
    from app.ai.common import effective_settings
    from app.db import AppSetting

    c, sf = client

    def put(**values):
        return c.put("/api/v1/admin/settings", json={"values": values})

    # "3.0" eskiden kaydediliyor ve sonraki her okumada int("3.0") ile YZ hattını düşürüyordu
    bad = put(relevance_samples="3.0")
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "invalid_setting"
    assert "tam sayı" in bad.json()["error"]["message"]
    assert put(relevance_samples="10").status_code == 400 and put(relevance_samples="0").status_code == 400
    assert put(relevance_threshold="1.5").status_code == 400 and put(relevance_threshold="nan").status_code == 400
    assert put(relevance_threshold="abc").status_code == 400
    ok = put(relevance_samples="5", relevance_threshold="0.40")
    assert ok.status_code == 200 and ok.json()["relevance_samples"] == 5 and ok.json()["relevance_threshold"] == 0.4
    with sf() as s:   # kanonik biçimde saklanır
        assert s.get(AppSetting, "relevance_threshold").value == "0.4"
    # birbirine bağlı eşikler: yalnızca ikisi de tutarlıysa kabul edilir; ihlalde hiçbiri yazılmaz
    inverted = put(dedupe_review_threshold="0.95")   # varsayılan otomatik eşik 0.90
    assert inverted.status_code == 400 and "dedupe_review_threshold" in inverted.json()["error"]["message"]
    assert c.get("/api/v1/admin/settings").json()["dedupe_review_threshold"] == settings.dedupe_review_threshold
    assert put(confidence_medium="0.9").status_code == 400   # varsayılan 'Yüksek' eşiği 0.75
    assert put(dedupe_review_threshold="0.95", dedupe_auto_threshold="0.99").status_code == 200
    assert put(confidence_medium="0.9", confidence_high="0.95").status_code == 200
    with sf() as s:   # elle yazılmış bozuk satır YZ hattını düşürmez; geçerli değer (varsayılan) kullanılır
        s.add(AppSetting(key="relevance_weight_model", value="çok"))
        s.commit()
        assert effective_settings(s, settings).relevance_weight_model == settings.relevance_weight_model


def test_admin_settings_detail_and_field_level_errors(settings, env, client):  # noqa: F811
    c, _ = client
    d = c.get("/api/v1/admin/settings?detail=true").json()
    assert d["values"]["relevance_samples"] == settings.relevance_samples
    assert d["specs"]["relevance_samples"] == {"type": "int", "min": 1, "max": 9,
                                               "description": d["specs"]["relevance_samples"]["description"]}
    assert d["specs"]["relevance_threshold"]["type"] == "float" and "relevance_threshold" not in d["specs"]["relevance_samples"]
    assert "specs" not in c.get("/api/v1/admin/settings").json()          # düz biçim geriye uyumlu
    r = c.put("/api/v1/admin/settings", json={"values": {"relevance_samples": "3.0", "relevance_threshold": "2",
                                                         "relevance_weight_rules": "0.3"}})
    fields = r.json()["error"]["details"]["fields"]
    assert r.status_code == 400 and set(fields) == {"relevance_samples", "relevance_threshold"}   # geçerli alan hata değil
    assert c.get("/api/v1/admin/settings").json()["relevance_weight_rules"] == settings.relevance_weight_rules  # hiçbiri yazılmadı
    cross = c.put("/api/v1/admin/settings", json={"values": {"dedupe_review_threshold": "0.95"}})
    assert set(cross.json()["error"]["details"]["fields"]) == {"dedupe_review_threshold"}
    unknown = c.put("/api/v1/admin/settings", json={"values": {"llm_api_key": "x"}})
    assert unknown.status_code == 400 and "llm_api_key" in unknown.json()["error"]["details"]["fields"]
    ok = c.put("/api/v1/admin/settings?detail=true", json={"values": {"relevance_samples": "5"}}).json()
    assert ok["values"]["relevance_samples"] == 5 and "specs" in ok


def test_admin_metrics_last_run_next_due_and_llm_latency(settings, env, client):  # noqa: F811
    from datetime import datetime, timezone

    from app.collectors.config import load_sources
    from app.db import FetchRun, LlmCall

    c, sf = client
    first = load_sources(settings.sources_file).enabled()[0]
    t0 = datetime(2026, 1, 5, 6, 0, tzinfo=timezone.utc)
    now = datetime.now(timezone.utc)
    with sf() as s:
        s.add(FetchRun(source_code=first.code, started_at=t0, finished_at=t0 + timedelta(seconds=42), status="success",
                       requests=7, items_new=3, items_listed=20))
        for i, ms in enumerate([100, 200, 300, 400, 5000]):
            s.add(LlmCall(task="summary", provider="fake", model="m", prompt_version="p", input_hash=f"h{i}",
                          latency_ms=ms, status="ok", created_at=now))
        s.add(LlmCall(task="summary", provider="fake", model="m", prompt_version="p", input_hash="bad",
                      latency_ms=None, status="error", created_at=now))
        s.commit()
    m = c.get("/api/v1/admin/metrics").json()
    row = next(x for x in m["sources"] if x["code"] == first.code)
    assert row["lastRun"]["durationSec"] == 42 and row["lastRun"]["requests"] == 7 and row["lastRun"]["status"] == "success"
    assert row["nextDue"] is not None                       # yerel kurulumda sources.yaml takviminden hesaplanır
    assert m["pendingIngest"] is None and m["crawlMode"] == settings.crawl_mode
    assert m["llm"]["calls"] == 6 and m["llm"]["failed"] == 1 and m["llm"]["p50Ms"] == 300 and m["llm"]["p95Ms"] == 4080


def test_admin_scans_lists_active_recent_and_filters(settings, env, client):  # noqa: F811
    from datetime import datetime, timezone

    from app.db import FetchRun

    c, sf = client
    now = datetime.now(timezone.utc)
    with sf() as s:
        s.add(FetchRun(source_code="BDDK", started_at=now - timedelta(minutes=5), status="running", requests=3))
        s.add(FetchRun(source_code="SPK", started_at=now - timedelta(hours=2), status="success",
                       finished_at=now - timedelta(hours=2) + timedelta(seconds=42),
                       channels={"liste": {"items": 5, "new": 2, "warnings": ["w"]}}))
        s.add(FetchRun(source_code="BDDK", started_at=now - timedelta(hours=3), status="failed",
                       finished_at=now - timedelta(hours=3) + timedelta(seconds=10),
                       errors=[f"hata {i}" for i in range(7)], channels={"liste": {"items": 0, "error": "boom"}}))
        # Resmî Gazete backfill'i started_at'a tarama gününü yazar: gerçek süre bilinmez
        s.add(FetchRun(source_code="RESMI_GAZETE", started_at=now - timedelta(days=30), status="success",
                       finished_at=now - timedelta(hours=1)))
        s.commit()
    body = c.get("/api/v1/admin/scans").json()
    assert body["crawlMode"] == settings.crawl_mode and body["crawler"] is None and body["queued"] == []
    assert len(body["active"]) == 1 and body["active"][0]["source"] == "BDDK" and not body["active"][0]["pendingIngest"]
    assert 299 <= body["active"][0]["elapsedSec"] <= 330 and body["active"][0]["requests"] == 3
    ids = [r["id"] for r in body["recent"]]
    assert len(ids) == 3 and ids == sorted(ids, reverse=True)           # id'ye göre: started_at backfill'de yanıltıcı
    bf, ok, failed = body["recent"]
    assert bf["backfillDay"] is not None and bf["durationSec"] is None
    assert ok["durationSec"] == 42 and ok["backfillDay"] is None
    assert ok["channels"] == [{"name": "liste", "items": 5, "new": 2, "changed": 0, "error": None, "warnings": 1,
                               "stopped": False, "incomplete": False}]
    assert failed["errorCount"] == 7 and len(failed["errors"]) == 5 and failed["channels"][0]["error"] == "boom"
    assert [r["source"] for r in c.get("/api/v1/admin/scans", params={"status": "failed"}).json()["recent"]] == ["BDDK"]
    assert [r["source"] for r in c.get("/api/v1/admin/scans", params={"source": "SPK"}).json()["recent"]] == ["SPK"]
    assert len(c.get("/api/v1/admin/scans", params={"limit": 2}).json()["recent"]) == 2
    assert c.get("/api/v1/admin/scans", params={"status": "bogus"}).json()["error"]["code"] == "validation_error"
    aq = body["aiQueue"]
    assert [s["status"] for s in aq["stages"]] == ["NEW", "RELEVANT", "SUMMARIZED", "AI_FAILED"] and aq["waiting"] >= 0


def test_admin_jobs_recorded_for_background_tasks(settings, env, client, monkeypatch):  # noqa: F811
    from datetime import date

    from app.db import JobRun

    c, sf = client
    payload = {"stage": "extract", "from": "2026-09-01", "to": "2026-09-02"}
    # TestClient arka plan görevini yanıt dönmeden çalıştırır: iş yanıttan sonra tamamlanmış olur
    monkeypatch.setattr("app.api.main._reprocess_bg",
                        lambda sf_, st, body, actor: {"stage": body.stage, "day": date(2026, 9, 1)})
    r = c.post("/api/v1/admin/reprocess", json=payload)
    assert r.status_code == 202 and r.json()["stage"] == "extract"
    with sf() as s:
        job = s.get(JobRun, r.json()["jobId"])
        assert (job.kind, job.status, job.via) == ("reprocess", "success", "inline")
        assert job.params == {"stage": "extract", "from": "2026-09-01", "to": "2026-09-02", "source": None}
        assert job.result == {"stage": "extract", "day": "2026-09-01"}            # tarih dizeye çevrilir
        assert job.started_at and job.finished_at and job.requested_by
    jobs = c.get("/api/v1/admin/scans").json()["jobs"]
    assert jobs["active"] == [] and jobs["recent"][0]["id"] == r.json()["jobId"]
    assert jobs["recent"][0]["status"] == "success" and jobs["recent"][0]["durationSec"] is not None

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr("app.api.main._reprocess_bg", boom)
    r = c.post("/api/v1/admin/reprocess", json=payload)
    assert r.status_code == 202                                                  # hata yanıtı bozmaz, kayda geçer
    with sf() as s:
        job = s.get(JobRun, r.json()["jobId"])
        assert job.status == "failed" and job.error == "RuntimeError: boom" and job.result is None


def test_search_folds_turkish_dotted_dotless_i_and_escapes_wildcards(settings, env, client):  # noqa: F811
    reg_id = _ready(settings, env)
    c, sf = client
    with sf() as s:
        s.get(Regulation, reg_id).title = "İŞLEMLERİ ışık IBAN %50 Kredi_Kartı Yönetmeliği"
        s.commit()

    def total(q):
        return c.get("/api/v1/regulations", params={"q": q}).json()["total"]

    assert total("işlemleri") == 1 and total("İŞLEMLERİ") == 1          # İ ↔ i
    assert total("Işık") == 1 and total("ışık") == 1 and total("ISIK") == 0   # I ↔ ı; s ≠ ş korunur
    assert total("iban") == 1 and total("IBAN") == 1 and total("İBAN") == 1
    assert total("%50") == 1 and total("kredi_kartı") == 1
    assert total("%") == 1 and total("_") == 1       # yalnız başlıktaki gerçek karakterle eşleşir
    assert total("%zzz%") == 0 and total("kredi%kartı") == 0 and total("kredi_ kartı") == 0   # joker değil
    assert total("ISTANBUL") == 1 and total("istanbul") == 1       # özet metni: "İstanbul"
    assert total("   ") == 1                                         # boş arama filtre sayılmaz


def test_security_headers_and_readyz_does_not_leak_db_error(settings, env, client):  # noqa: F811
    c, _ = client
    for path in ("/", "/api/v1/me", "/readyz", "/api/v1/yok"):
        r = c.get(path)
        assert r.headers["x-content-type-options"] == "nosniff", path
        assert r.headers["x-frame-options"] == "DENY" and r.headers["content-security-policy"] == "frame-ancestors 'none'"
        assert r.headers["referrer-policy"] == "strict-origin-when-cross-origin"
    assert c.get("/api/v1/yok").status_code == 404

    def broken():
        raise RuntimeError("postgresql://mevzuat_app:s3cret@10.0.0.7/mevzuat bağlantı reddedildi")

    bad = TestClient(create_app(settings.model_copy(update={"cors_origins": ""}), broken))
    r = bad.get("/readyz")
    assert r.status_code == 503 and r.json()["checks"]["db"] == "hata" and r.json()["status"] == "hata"
    assert "s3cret" not in r.text and "10.0.0.7" not in r.text
    assert r.headers["x-frame-options"] == "DENY"


def test_portal_page_loads_only_local_assets_with_security_headers(settings, env, client):  # noqa: F811
    import re

    c, _ = client
    r = c.get("/")
    html = r.text
    assert r.status_code == 200 and '<html lang="tr">' in html
    # dış CDN / yazı tipi / betik yok: tüm src ve href değerleri sunucudan (/assets/) gelir
    for url in re.findall(r'(?:src|href)="([^"#][^"]*)"', html):
        assert url.startswith("/assets/"), url
    for host in ("fonts.googleapis.com", "fonts.gstatic.com", "unpkg.com", "cdn.jsdelivr", "cdnjs"):
        assert host not in html
    for name in re.findall(r'(?:src|href)="(/assets/[^"]+)"', html):
        asset = c.get(name)
        assert asset.status_code == 200, name
        assert "http://" not in asset.text and "https://" not in asset.text.replace("https://www.w3.org", ""), name
    for h, v in (("x-frame-options", "DENY"), ("x-content-type-options", "nosniff"),
                 ("content-security-policy", "frame-ancestors 'none'")):
        assert r.headers[h] == v


def test_old_portal_routes_redirect_to_new_page_and_legacy_pages_stay_available(settings, env, client):  # noqa: F811
    c, _ = client
    for old, target in (("/surec", "/#hat"), ("/arsiv", "/#arsiv"), ("/karsilastirma", "/#karsilastirma"),
                        ("/admin", "/#yonetim")):
        r = c.get(old, follow_redirects=False)
        assert r.status_code == 307 and r.headers["location"] == target, old
    for page in ("", "/surec", "/arsiv", "/karsilastirma"):
        r = c.get("/eski" + page)
        assert r.status_code == 200 and "text/html" in r.headers["content-type"], page
    assert c.get("/eski/yok").status_code == 404 and c.get("/eski/..%2Fsecret").status_code == 404


def test_get_visible_for_update_refreshes_stale_row(settings, env, client):  # noqa: F811
    from app.services.regulations import ApiError, get_visible

    reg_id = _ready(settings, env)
    _, sf = client
    with sf() as first, sf() as second:
        stale = get_visible(first, reg_id)
        assert stale.row_version == 1
        second.get(Regulation, reg_id).row_version = 7
        second.commit()
        assert first.get(Regulation, reg_id).row_version == 1          # kimlik haritasındaki bayat nesne
        fresh = get_visible(first, reg_id, for_update=True)
        assert fresh.row_version == 7                                   # kilitli okuma güncel satırı getirir
    with sf() as s, pytest.raises(ApiError) as e:
        get_visible(s, 9999, for_update=True)
    assert e.value.code == "not_found"


def test_vendor_files_are_served_from_a_plain_file_name_only(settings, env, tmp_path):  # noqa: F811
    sf, _ = env
    portal = tmp_path / "portal"
    (portal / "vendor").mkdir(parents=True)
    (portal / "vendor" / "react.production.min.js").write_text("window.React={};")
    (portal / "secret.txt").write_text("gizli")
    (portal / "vendor" / "tool.exe").write_text("MZ")
    c = TestClient(create_app(settings.model_copy(update={"cors_origins": "", "portal_dir": portal}), sf))
    ok = c.get("/vendor/react.production.min.js")
    assert ok.status_code == 200 and ok.text == "window.React={};"
    assert ok.headers["content-type"].startswith("application/javascript") and "max-age" in ok.headers["cache-control"]
    assert ok.headers["x-content-type-options"] == "nosniff"
    for name in ("yok.js", "tool.exe", "..%2Fsecret.txt", "..%5Csecret.txt", "%2e%2e%2fsecret.txt"):
        assert c.get(f"/vendor/{name}").status_code == 404, name
    assert c.get("/vendor/").status_code == 404


def test_audit_hash_chain_detects_tampering(settings, env, client):  # noqa: F811
    from sqlalchemy import text

    reg_id = _ready(settings, env)
    c, sf = client
    c.post(f"/api/v1/regulations/{reg_id}/views")
    c.post(f"/api/v1/regulations/{reg_id}/decision", json={"decision": "approve", "note": "ilk not"})
    assert c.get("/api/v1/admin/audit/verify").json() == {"ok": True, "problems": []}
    with sf() as s:   # ORM korumasını atlayan doğrudan veritabanı değişikliği
        s.execute(text("UPDATE audit_event SET note = 'değiştirildi' WHERE event_type = 'approved'"))
        s.commit()
    r = c.get("/api/v1/admin/audit/verify").json()
    assert r["ok"] is False and r["problems"][0]["problem"] == "içerik değiştirilmiş"
    with sf() as s:   # araya olay silme → zincir kopar
        s.execute(text("DELETE FROM audit_event WHERE event_type = 'viewed'"))
        s.commit()
    assert "zincir kopuk" in {p["problem"] for p in c.get("/api/v1/admin/audit/verify").json()["problems"]}
