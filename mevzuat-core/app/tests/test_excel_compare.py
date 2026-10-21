"""Başkanlık Excel listesi ↔ sistem kayıtları karşılaştırması (/api/v1/evaluation/excel-compare, salt okunur)."""
from __future__ import annotations

from datetime import date
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.api.main import create_app
from app.db import ManualDetection, RawDocument, Regulation
from app.evaluation.excel_compare import build_workbook, clean_body, compare, parse_sheet
from tests.test_processing import add_raw, env  # noqa: F401

HEADERS = ["Referans No", "Konu", "Talep Tipi", "Gönderen/Düzenleyici Kurum", "Risk Seviyesi", "Alıcılar",
           "Oluşturulma Tarihi", "Metin"]
BDDK = "Bankacılık Düzenleme ve Denetleme Kurumu"
BODY = ("Kamuoyunda Fon Soruşturması olarak bilinen süreç ile bağlantılı olarak son dönemde yaşanan gelişmeler "
        "kapsamında Kurumumuzca yapılan değerlendirmeler sonucunda Bankacılık Düzenleme ve Denetleme Kurulunun "
        "30.09.2026 tarihli ve 11576, 11577 sayılı Kararları ile 5411 sayılı Bankacılık Kanununun 18 inci maddesinin "
        "beşinci fıkrası uyarınca Destek Yatırım Bankası A.Ş.'de nitelikli paya sahip olan ortaklara ait paylara "
        "ilişkin temettü dışındaki ortaklık haklarının Tasarruf Mevduatı Sigorta Fonu tarafından kullanılmasına "
        "karar verilmiştir.")
# Excel'deki biçim: giriş cümlesi, &nbsp; varlıkları, "****" ayracı, satır içi kırılmalar
METIN = ("Bankacılık Düzenleme ve Denetleme Kurumu tarafından 30.09.2026 tarihinde yayınlanan Basın Duyurusu ektedir."
         "Basın Duyurusu:&nbsp;Duyuru Detay************************************************************"
         + BODY.replace("son dönemde ", "son dönemde \n").replace(" ile 5411", " ile&nbsp;&nbsp;5411"))
OTHER_BODY = ("Kurumumuz internet sitesinde yer alan bankacılık sektörü aylık bülteninin yayım takvimi güncellenmiş "
              "olup kamuoyunun bilgisine sunulur. Bülten verileri her ayın son iş günü yayımlanmaya devam edecektir.")


def _xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


EXCEL = _xlsx([
    ["MVZ-1", "Hedef Yatırım Bankası A.Ş. Hakkında Kurul Kararı", "Bilgilendirme", BDDK, "Yüksek", "Hazine",
     "01.10.2026 11:06", "Kurul kararı ektedir."],
    ["MVZ-2", "Destek Yatırım Bankası A.Ş., Hedef Yatırım Bankası A.Ş. ile Tera Finans Faktoring A.Ş. Hakkında "
              "Basın Duyurusu", "Bilgilendirme", BDDK, "Yüksek", "Yatırımcı İlişkileri; Fonlama",
     "01.10.2026 11:06", METIN],
    ["MVZ-3", "Elektronik tebligat usulüne ilişkin genelge", "Aksiyon", "Gelir İdaresi Başkanlığı", "Orta", "Vergi",
     "02.10.2026 09:00", "Elektronik tebligat sistemine ilişkin usul ve esaslar belirlenmiştir."],
])


def _reg(sf, key, title, status, pub, *, issuer="BDDK", relevant=None, severity=None, rationale=None) -> int:
    with sf() as s:
        reg = Regulation(title=title, issuer=issuer, issuer_name=issuer, primary_source_code=issuer,
                         canonical_key=key, processing_status=status, publish_date=pub, is_relevant=relevant,
                         relevance_score=0.8 if relevant else 0.1, severity=severity,
                         classification={"relevance": {"rationale": rationale}} if rationale else {})
        s.add(reg)
        s.commit()
        return reg.id


def _raw(sf, storage, reg_id, ext, title, text, pub):
    rid = add_raw(sf, storage, source="BDDK", channel="duyuru", external_id=ext, url=f"https://bddk.org.tr/{ext}",
                  title=title, published=pub)
    with sf() as s:
        raw = s.get(RawDocument, rid)
        raw.text, raw.regulation_id = text, reg_id
        s.commit()


@pytest.fixture
def data(settings, env):  # noqa: F811
    sf, storage = env
    ids = {
        "A": _reg(sf, "a", "Hedef Yatırım Bankası A.Ş. Hakkında Kurul Kararı", "READY", date(2026, 9, 30),
                  relevant=True, severity="Yüksek"),
        "B": _reg(sf, "b", "Basın Duyurusu", "IRRELEVANT", date(2026, 9, 30), relevant=False,
                  rationale="Genel kamuoyu duyurusu."),
        "B2": _reg(sf, "b2", "Basın Duyurusu", "READY", date(2026, 9, 29), relevant=True),
        "D": _reg(sf, "d", "Yatırım fonlarına ilişkin rehber", "READY", date(2026, 10, 1), issuer="SPK",
                  relevant=True),
        "E": _reg(sf, "e", "Eski düzenleme", "READY", date(2026, 8, 1), relevant=True),
    }
    _raw(sf, storage, ids["B2"], "b2", "Basın Duyurusu", OTHER_BODY, date(2026, 9, 29))
    _raw(sf, storage, ids["B"], "b", "Basın Duyurusu", BODY, date(2026, 9, 30))
    return sf, ids


def test_clean_body_strips_entities_and_intro():
    t = clean_body(METIN)
    assert t.startswith("Kamuoyunda Fon Soruşturması") and "&nbsp;" not in t and "ektedir" not in t
    assert clean_body("<p>Merhaba&nbsp;<b>dünya</b></p>") == "Merhaba dünya"


def test_compare_categories_text_fallback_and_reverse_direction(data):
    sf, ids = data
    rows = parse_sheet("liste.xlsx", EXCEL)
    assert [r["ref"] for r in rows] == ["MVZ-1", "MVZ-2", "MVZ-3"] and rows[0]["created"] == date(2026, 10, 1)
    with sf() as s:
        res = compare(s, rows)
    by_ref = {i["ref"]: i for i in res["items"]}
    # Excel tarihi (01.10) yayımdan (30.09) sonra: geniş pencereyle başlıktan bulunur
    a = by_ref["MVZ-1"]
    assert a["group"] == "relevant" and a["regulation"]["id"] == ids["A"] and a["method"] == "fuzzy_title"
    # genel "Basın Duyurusu" başlıklı iki aday: metin doğru (ilgisiz bulunmuş) kaydı seçer
    b = by_ref["MVZ-2"]
    assert b["group"] == "irrelevant" and b["regulation"]["id"] == ids["B"] and b["method"] == "fuzzy_text"
    assert b["regulation"]["rationale"] == "Genel kamuoyu duyurusu."
    c = by_ref["MVZ-3"]
    assert c["group"] == "missing" and c["category"] == "not_collected" and c["regulation"] is None

    assert res["summary"] == {"total": 3, "relevant": 1, "irrelevant": 1, "missing": 1,
                              "missingByReason": {"not_collected": 1}, "recall": 0.333, "foundRate": 0.667}
    # ters yön: 28.09–02.10 arasında görünen A, B2, D; listede yalnız A var. E dönem dışı.
    p = res["precision"]
    assert p["period"] == {"from": "2026-09-28", "to": "2026-10-02"}
    assert p["shown"] == 3 and p["inList"] == 1 and p["precision"] == 0.333
    assert {r["id"] for r in res["extra"]} == {ids["B2"], ids["D"]}
    assert res["severityAgreement"] == {"n": 1, "exact": 1.0}

    wb = openpyxl.load_workbook(BytesIO(build_workbook(res)))
    assert wb.sheetnames == ["Özet", "Excel satırları", "Fazla uyarılar"]
    assert wb["Excel satırları"].max_row == 4 and wb["Fazla uyarılar"].max_row == 3


def test_title_threshold_parameter(data):
    sf, ids = data
    rows = parse_sheet("liste.xlsx", EXCEL)[:1]
    rows[0]["title"] = "Hedef Yatırım Bankası Kurul Kararı hakkında bilgi notu"
    rows[0]["text"] = ""
    with sf() as s:
        assert compare(s, rows, title_threshold=100)["items"][0]["group"] == "missing"
        assert compare(s, rows, title_threshold=60)["items"][0]["regulation"]["id"] == ids["A"]


def test_api_json_and_xlsx_do_not_write(settings, data):
    sf, ids = data
    c = TestClient(create_app(settings.model_copy(update={"cors_origins": ""}), sf))
    files = {"file": ("liste.xlsx", EXCEL, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = c.post("/api/v1/evaluation/excel-compare", files=files)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["total"] == 3
    x = c.post("/api/v1/evaluation/excel-compare", params={"format": "xlsx"}, files=files)
    assert x.status_code == 200 and 'filename="karsilastirma-' in x.headers["content-disposition"]
    assert openpyxl.load_workbook(BytesIO(x.content)).sheetnames[0] == "Özet"
    with sf() as s:
        assert s.scalar(select(func.count()).select_from(ManualDetection)) == 0
    bad = c.post("/api/v1/evaluation/excel-compare", files={"file": ("liste.pdf", b"x", "application/pdf")})
    assert bad.json()["error"]["code"] == "unsupported_file"
    empty = c.post("/api/v1/evaluation/excel-compare", files={"file": ("bos.xlsx", _xlsx([]), "application/xlsx")})
    assert empty.json()["error"]["code"] == "empty_file"
    assert c.get("/karsilastirma", follow_redirects=False).headers["location"] == "/#karsilastirma"
    assert "Excel karşılaştırma" in c.get("/assets/views.js").text   # ekran yeni portalın betiğinde
