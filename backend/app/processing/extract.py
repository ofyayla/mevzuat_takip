"""Ham içerikten metin çıkarma (İK-2): HTML, PDF (+ OCR), DOCX, görüntü, liste satırı JSON'u.

- HTML: kanalın ``content_selector``'ı tanımlıysa o bölüm (ör. BDDK detay sayfasında trafilatura çerez
  penceresini ana metin sanıyor); yoksa trafilatura ile ana içerik (tablolar Markdown). trafilatura boş dönerse
  menüsüz gövde metnine düşülür.
- PDF: PyMuPDF ile sayfa sayfa metin katmanı. Metni ``ocr_min_chars_per_page`` altında kalan sayfalar
  (taranmış) OCR servisine yalnızca o sayfalar olarak gönderilir; sayfalar ``\\f`` ile ayrılır.
- OCR yapılandırılmamışsa veya belge ilk taramaya aitse (``ocr_baseline=False``) taranmış sayfalar boş kalır ve
  ``ocr_pending_pages`` doldurulur; belge daha sonra ``mevzuat-process reprocess`` ile yeniden işlenebilir.
"""
from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass, field

from app.collectors.http import Response
from app.collectors.strategies._html import clean, parse_html
from app.processing.normalize import normalize_text, text_quality
from app.processing.ocr import OcrEngine
from app.settings import Settings

log = logging.getLogger(__name__)

PAGE_SEP = "\f"
# Taranmış sayfa: metin katmanı yok denecek kadar az, ya da kısa bir üst bilgi + sayfayı kaplayan görüntü
# (RG'nin taranmış PDF'lerinde ilk sayfada ~120 karakterlik "Resmî Gazete Sayı : …" metni bulunuyor)
SCANNED_HEADER_MAX_CHARS = 250
SCANNED_MIN_IMAGE_COVERAGE = 0.25
MIN_TEXT_QUALITY = 0.6
IMAGE_TYPES = {"image/jpeg", "image/png", "image/tiff", "image/bmp", "image/heif"}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class UnsupportedContent(Exception):
    pass


@dataclass
class ExtractResult:
    text: str
    method: str
    page_count: int | None = None
    ocr_confidence: float | None = None
    ocr_pages: list[int] = field(default_factory=list)
    ocr_pending_pages: list[int] = field(default_factory=list)


def decode_html(content: bytes) -> str:
    return Response("", "", 200, {"content-type": "text/html"}, content).text


_BLOCK_TAG = re.compile(r"<(?:br|/?p|/?div|/?tr|/?li|/?h[1-6]|/?table|/?ul|/?ol|/?section|/?article|/?blockquote)\b[^>]*>",
                        re.I)
_CELL_END = re.compile(r"</t[dh]\s*>", re.I)


def html_to_text(html_text: str, selector: str | None = None) -> str:
    """Blok elemanlarında satır kıran HTML→metin. Kaynaktaki satır sarmaları (Word HTML'inde cümle ortasında
    yeni satır) boşluğa çevrilir; paragraf ve tablo satırı sınırları korunur."""
    html_text = re.sub(r"[\r\n]+", " ", html_text)
    html_text = _BLOCK_TAG.sub(lambda m: "\n" + m.group(0), html_text)
    html_text = _CELL_END.sub(lambda m: m.group(0) + " ", html_text)
    tree = parse_html(html_text)
    node = tree.css_first(selector) if selector else tree.body
    return normalize_text(node.text(separator="")) if node is not None else ""


def extract_html(content: bytes, content_selector: str | None = None) -> ExtractResult:
    """Önce kanalın ``content_selector``'ı (varsa), sonra trafilatura; ikisi de boşsa menüsüz gövde metni."""
    html_text = decode_html(content)
    if content_selector and (text := html_to_text(html_text, content_selector)):
        return ExtractResult(text, "html_selector")
    import trafilatura

    main = trafilatura.extract(html_text, output_format="markdown", include_tables=True, include_comments=False,
                               favor_recall=True, deduplicate=False) or ""
    main = normalize_text(main)
    if len(main) >= 50:
        return ExtractResult(main, "html")
    fallback = html_to_text(html_text)
    return ExtractResult(fallback if len(fallback) > len(main) else main, "html_body")


def extract_pdf(content: bytes, settings: Settings, ocr: OcrEngine | None, *, allow_ocr: bool) -> ExtractResult:
    import pymupdf

    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
    except Exception as e:  # noqa: BLE001 — bozuk/şifreli PDF
        raise UnsupportedContent(f"PDF açılamadı: {e}") from e
    with doc:
        if doc.needs_pass:
            raise UnsupportedContent("PDF parola korumalı")
        texts = [_page_text(page) for page in doc]
        scanned = [i + 1 for i, page in enumerate(doc) if _needs_ocr(page, texts[i], settings)]
    res = ExtractResult("", "pdf_text", page_count=len(texts))
    if scanned and ocr is not None and allow_ocr:
        o = ocr.analyze(content, "application/pdf", pages=scanned)
        for n in scanned:
            if (t := o.pages.get(n)) and len(t.strip()) > len(texts[n - 1].strip()):
                texts[n - 1] = t
        res.ocr_pages, res.ocr_confidence = scanned, o.confidence
        res.method = "ocr" if len(scanned) == len(texts) else "pdf_text+ocr"
    elif scanned:
        res.ocr_pending_pages = scanned
    res.text = PAGE_SEP.join(normalize_text(t) for t in texts)
    return res


def _page_text(page) -> str:
    text = page.get_text("text")
    return text if text_quality(text) >= MIN_TEXT_QUALITY else ""   # çöp metin katmanı = metin yok


def _needs_ocr(page, text: str, settings: Settings) -> bool:
    import pymupdf

    chars = len(text.strip())
    if chars < settings.ocr_min_chars_per_page:
        return True
    if chars >= SCANNED_HEADER_MAX_CHARS:
        return False
    area = abs(page.rect) or 1
    covered = sum(abs(pymupdf.Rect(i["bbox"]) & page.rect) for i in page.get_image_info())
    return covered / area >= SCANNED_MIN_IMAGE_COVERAGE


def extract_image(content: bytes, content_type: str, ocr: OcrEngine | None, *, allow_ocr: bool) -> ExtractResult:
    if ocr is None or not allow_ocr:
        return ExtractResult("", "image", page_count=1, ocr_pending_pages=[1])
    o = ocr.analyze(content, content_type)
    text = PAGE_SEP.join(normalize_text(o.pages[n]) for n in sorted(o.pages))
    return ExtractResult(text, "ocr", page_count=len(o.pages) or 1, ocr_confidence=o.confidence,
                         ocr_pages=sorted(o.pages))


def extract_docx(content: bytes) -> ExtractResult:
    import docx

    try:
        d = docx.Document(io.BytesIO(content))
    except Exception as e:  # noqa: BLE001
        raise UnsupportedContent(f"DOCX açılamadı: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("| " + " | ".join(clean(c.text) for c in row.cells) + " |")
    return ExtractResult(normalize_text("\n".join(parts)), "docx")


def extract_listing(content: bytes) -> ExtractResult:
    """Dış siteye giden öğe için saklanan liste satırı: başlık + özet + ek alanlar."""
    data = json.loads(content)
    parts = [data.get("title") or "", data.get("summary") or ""]
    parts += [f"{k}: {v}" for k, v in (data.get("extra") or {}).items() if isinstance(v, (str, int, float))]
    return ExtractResult(normalize_text("\n".join(p for p in parts if p)), "listing")


def extract(content: bytes, content_type: str, role: str, settings: Settings, ocr: OcrEngine | None, *,
            allow_ocr: bool = True, content_selector: str | None = None) -> ExtractResult:
    ct = (content_type or "").split(";")[0].strip().lower()
    if role == "listing" or (ct == "application/json" and content[:1] == b"{" and b'"title"' in content[:200]):
        return extract_listing(content)
    if ct == "application/pdf" or content[:5] == b"%PDF-":
        return extract_pdf(content, settings, ocr, allow_ocr=allow_ocr)
    if ct in ("text/html", "application/xhtml+xml") or content[:200].lstrip().lower().startswith((b"<!doctype", b"<html")):
        return extract_html(content, content_selector)
    if ct == DOCX_TYPE:
        return extract_docx(content)
    if ct in IMAGE_TYPES:
        return extract_image(content, ct, ocr, allow_ocr=allow_ocr)
    if ct.startswith("text/"):
        return ExtractResult(normalize_text(content.decode("utf-8", errors="replace")), "text")
    raise UnsupportedContent(f"desteklenmeyen içerik türü: {ct or 'bilinmiyor'}")
